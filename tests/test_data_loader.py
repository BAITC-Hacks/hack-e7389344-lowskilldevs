"""Synthetic Excel fixtures exercise the supplier layouts without partner data."""
from datetime import datetime
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import Workbook

from data_loader import combine_datasets, load_partner_files


def workbook(rows, title="Лист_1"):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    wb.save(output)
    return output.getvalue()


class DataLoaderTests(unittest.TestCase):
    def test_sparse_monthly_sales_keep_code_and_exclude_totals(self):
        payload = workbook([
            ["Номенклатура", "Номенклатура.Код", "янв. 2026", "февр. 2026", "Итого"],
            [None, None, "Количество", "Количество", "Количество"],
            ["Тест", "000123_", 10, None, 10],
            ["Возвраты", "000456_", -2, "3,5", 1.5],
            ["Итого", None, 8, 3.5, 11.5],
            ["Группа", None, 99, 99, 198],
        ])
        data = load_partner_files({"Ежемесячные продажи.xlsx": payload}, "IEK")
        sales = data["sales"]
        self.assertEqual(len(sales), 4)
        self.assertEqual(set(sales.sku), {"000123_", "000456_"})
        self.assertEqual(sales.qty.sum(), 11.5)
        self.assertEqual(sales.iloc[1].qty, 0)
        self.assertEqual(sales.iloc[0].month, pd.Timestamp("2026-01-01"))

    def test_numeric_codes_honor_leading_zero_excel_format(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["sku", "month", "qty"])
        ws.append([123, "2026-01", 5])
        ws["A2"].number_format = "000000"
        stream = BytesIO()
        wb.save(stream)
        data = load_partner_files({"export.xlsx": stream.getvalue()}, "S")
        self.assertEqual(data["sales"].iloc[0].sku, "000123")

    def test_opening_inventory_blank_is_unknown(self):
        data = load_partner_files({"Ежемесячные остатки.xlsx": workbook([
            ["Номенклатура", "Ед.", "Номенклатура.Код", "авг. 2026", "сент. 2026", "Итого"],
            [None, None, None, "Количество", "Количество", "Количество"],
            [None, None, None, "нач. остаток", "нач. остаток", "нач. остаток"],
            ["Тест", "шт", "01", 5, None, 5],
        ])}, "S")
        self.assertEqual(len(data["stock"]), 2)
        self.assertTrue(pd.isna(data["stock"].iloc[1].stock))
        self.assertEqual(set(data["stock"].stock_basis), {"opening"})

    def test_moq_se_and_iek_have_different_semantics(self):
        se = load_partner_files({"MOQ.xlsx": workbook([
            ["№", "Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"],
            [1, "Тест", "01", "EXT-1", 6],
        ])}, "SE")
        iek = load_partner_files({"MOQ.xlsx": workbook([
            ["№", "Код 1с", "Артикул поставщика", "Наименование", "Мин. разр. к отгр."],
            [1, "01", "EXT-1", "Тест", 5],
        ])}, "IEK")
        self.assertEqual(tuple(se["moq"].iloc[0][["moq", "pack_size"]]), (6, 6))
        self.assertEqual(tuple(iek["moq"].iloc[0][["moq", "pack_size"]]), (5, 1))
        combined = combine_datasets([se, iek])
        self.assertEqual(len(combined["moq"]), 2)

    def test_iek_transit_order_dates_and_quantities(self):
        data = load_partner_files({"Путь ИЭК.xlsx": workbook([
            ["Код 1с", "Артикул ИЭК", "Наименование", "Заказ от 01.09.2026 (поступление до 10.10.2026)", "Заказ (поступление до 30.09.2026)"],
            ["0001", "EXT1", "Тест", 10, 20],
            ["0002", "EXT2", "Тест2", None, 0],
        ])}, "IEK")
        transit = data["transit"]
        self.assertEqual(transit.qty.sum(), 30)
        self.assertEqual(set(transit.eta), {pd.Timestamp("2026-10-10"), pd.Timestamp("2026-09-30")})

    def test_se_planning_transit_does_not_duplicate_sales(self):
        data = load_partner_files({"Товар в пути_SE на 22.09.2026.xlsx": workbook([
            [None, None, "СКЛАДЫ"],
            ["№", "Артикул поставщика", "Код 1с", "Наименование", "Категория 2026", "янв. 2026", "Свободный остаток", "СЭ в пути 24.09"],
            [1, "EXT1", "0001", "Тест", "A", 999, 15, 20],
        ]), "sales.xlsx": workbook([
            ["Номенклатура", "Номенклатура.Код", "янв. 2026"],
            ["Тест", "0001", 10],
        ])}, "SE")
        self.assertEqual(data["sales"].qty.sum(), 10)
        self.assertEqual(data["sales"].iloc[0].category, "A")
        self.assertEqual(data["transit"].qty.sum(), 20)
        self.assertTrue(data["transit"].eta.isna().all())
        self.assertEqual(data["stock"].iloc[0].month, pd.Timestamp("2026-09-22"))
        self.assertEqual(data["stock"].iloc[0].stock_basis, "snapshot")

    def test_seasonality_uses_final_factor_not_revenue_or_total(self):
        data = load_partner_files({"Сезонность.xlsx": workbook([
            [], [], ["год", "янв", "фев"], [2025, 1000, 2000], [],
            [None, "Месяц", "Продажи 2025", "Коэф. сезонности", "СЕЗОННОСТЬ"],
            [None, "янв", 1000, .6, .8],
            [None, "фев", 2000, 1.4, 1.2],
            [None, "Итого", 3000, 1, 1],
        ])}, "S")
        self.assertEqual(list(data["seasonality"].factor), [.8, 1.2])
        self.assertEqual(list(data["seasonality"].month), [1, 2])

    def test_journal_growth_uses_matched_complete_months_without_double_count(self):
        tx = workbook([
            ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"],
            ["01.01.2025 10:00:00", "d1", "Private document", "001", "Тест", "шт", "Склад", 100],
            ["01.01.2026 10:00:00", "d2", "Private document", "001", "Тест", "шт", "Склад", 160],
            ["02.01.2026 10:00:00", "d3", "Private document", "001", "Тест", "шт", "Склад", -10],
            ["05.02.2026 10:00:00", "d4", "Private document", "001", "Тест", "шт", "Склад", 9999],
        ])
        monthly = workbook([["sku", "month", "qty"], ["001", "2026-01", 150]])
        data = load_partner_files({"Динамика продаж.xlsx": tx, "sales.xlsx": monthly}, "S")
        self.assertEqual(data["sales"].qty.sum(), 150)
        self.assertEqual(data["growth"].iloc[0].growth_rate, .5)
        self.assertEqual(len(data["transactions"]), 4)
        self.assertEqual(set(data["transactions"].customer_id), {""})
        self.assertNotIn("Private document", str(data))

    def test_journal_fallback_and_customer_hashing(self):
        data = load_partner_files({"export.xlsx": workbook([
            ["date", "sku", "qty", "клиент"],
            [datetime(2026, 1, 1), "001", 10, "Sensitive person"],
            [datetime(2026, 1, 2), "001", 2, "Sensitive person"],
        ])}, "S")
        self.assertEqual(data["sales"].qty.sum(), 12)
        self.assertEqual(data["transactions"].customer_id.nunique(), 1)
        self.assertNotIn("Sensitive person", str(data))

    def test_malformed_duplicate_and_invalid_values_are_reported(self):
        valid = workbook([["sku", "month", "qty"], ["001", "2026-01", 10], ["002", "bad date", 99], ["003", "2026-01", "not a number"]])
        data = load_partner_files({"broken.xlsx": b"this is not a workbook", "one.xlsx": valid, "two.xlsx": valid}, "S")
        self.assertEqual(data["sales"].qty.sum(), 10)
        self.assertTrue(any("не прочитан" in w for w in data["warnings"]))
        self.assertTrue(any("повторная копия" in w for w in data["warnings"]))
        self.assertTrue(any("неверный месяц" in w for w in data["warnings"]))


if __name__ == "__main__":
    unittest.main()
