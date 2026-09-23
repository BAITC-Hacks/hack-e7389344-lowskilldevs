"""Read the partner's 1C Excel exports locally into a small, explicit schema.

The monthly quantity export is authoritative. The detailed sales journal is kept
separately for anomaly analysis and comparable-period growth; it is never added
to an existing monthly quantity series. Customer names and document text are
never returned. Reading a workbook does not execute its formulas or macros.
"""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from itertools import chain, islice
from pathlib import Path
import math
import re
import secrets
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

SCHEMAS = {
    "sales": ["supplier", "sku", "name", "category", "month", "qty"],
    "stock": ["supplier", "sku", "month", "stock", "stock_basis"],
    "moq": ["supplier", "sku", "moq", "pack_size"],
    "transit": ["supplier", "sku", "qty", "eta"],
    "seasonality": ["supplier", "category", "month", "factor"],
    "growth": ["supplier", "category", "growth_rate"],
    "transactions": ["supplier", "sku", "date", "qty", "customer_id"],
}
ALL_CATEGORY = "Все"
_MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5,
    "июн": 6, "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value).lower().replace("ё", "е"))


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
        if not value:
            return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def _sku(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    if value is None or isinstance(value, (bool, datetime, date)):
        return ""
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 0 or value != int(value):
            return ""
        value = str(int(value))
        number_format = getattr(cell, "number_format", "")
        if re.fullmatch(r"0+", number_format):
            value = value.zfill(len(number_format))
    value = str(value).strip()
    if _norm(value) in {"итого", "всего", "total", "nan", "none", "код", "артикул", "sku"}:
        return ""
    return value


def _date(value: Any) -> pd.Timestamp | None:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        result = pd.Timestamp(value)
        return None if pd.isna(result) else result.tz_localize(None)
    if isinstance(value, (int, float)) and 20000 < value < 100000:
        return pd.Timestamp(from_excel(value))
    value = _text(value)
    for pattern in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y"):
        try:
            return pd.Timestamp(datetime.strptime(value, pattern))
        except ValueError:
            pass
    return None


def _month_number(value: Any) -> int | None:
    number = _number(value)
    if number is not None and number == int(number) and 1 <= number <= 12:
        return int(number)
    return _MONTHS.get(_norm(value)[:3])


def _month(value: Any) -> pd.Timestamp | None:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(value).to_period("M").to_timestamp()
    text = _norm(value)
    year = re.search(r"\b(20\d{2})\b", text)
    number = _month_number(text)
    if year and number:
        return pd.Timestamp(int(year.group()), number, 1)
    if re.fullmatch(r"20\d{2}-\d{1,2}(?:-\d{1,2})?", text):
        try:
            return pd.Timestamp(text).to_period("M").to_timestamp()
        except ValueError:
            return None
    return None


def _find(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for candidate in candidates:
        if candidate in headers:
            return headers.index(candidate)
    return None


def _column_map(headers: list[str]) -> dict[str, int | None]:
    return {
        "sku": _find(headers, ("номенклатура.код", "код 1с", "код 1c", "код", "sku", "артикул", "артикул поставщика")),
        "name": _find(headers, ("номенклатура", "наименование", "name", "товар")),
        "category": _find(headers, ("категория 2026", "категория", "category", "группа")),
        "date": _find(headers, ("дата", "date", "дата продажи")),
        "month": _find(headers, ("месяц", "month", "период")),
        "qty": _find(headers, ("количество", "qty", "кол-во", "продажи")),
        "stock": _find(headers, ("stock", "остаток", "свободный остаток")),
        "moq": _find(headers, ("moq", "минимальный заказ", "мин. разр. к отгр.", "минимальная партия")),
        "pack": _find(headers, ("pack_size", "кратность", "упаковка", "кратность упаковки")),
        "eta": _find(headers, ("eta", "дата поступления", "дата прибытия", "поступление")),
        "customer": _find(headers, ("customer_id", "клиент", "покупатель", "контрагент")),
        "factor": _find(headers, ("factor", "сезонность", "коэффициент сезонности", "коэф. сезонности")),
        "growth": _find(headers, ("growth_rate", "темп роста", "рост")),
    }


def _file_kind(filename: str) -> str | None:
    name = _norm(filename)
    if "moq" in name or "кратност" in name:
        return "moq"
    if "динамик" in name or "transaction" in name:
        return "transactions"
    if "сезон" in name or "season" in name:
        return "seasonality"
    if "путь" in name or "пути" in name or "transit" in name:
        return "transit"
    if "остат" in name or "stock" in name:
        return "stock"
    if "продаж" in name or "sales" in name:
        return "sales"
    if "growth" in name:
        return "growth"
    return None


def _detect(preview: list[tuple], preferred: str | None):
    """Locate data headers past report titles and repeated 1C subheaders."""
    candidates = []
    for index, row in enumerate(preview):
        values = [getattr(c, "value", c) for c in row]
        headers = [_norm(v) for v in values]
        cols = _column_map(headers)
        months = {i: _month(v) for i, v in enumerate(values) if _month(v) is not None}
        transit_columns = {i: v for i, v in enumerate(values) if "в пути" in _norm(v) or "поступление до" in _norm(v)}
        kind = None
        if cols["sku"] is not None:
            if transit_columns or (cols["eta"] is not None and cols["qty"] is not None):
                kind = "transit"
            elif cols["date"] is not None and cols["qty"] is not None:
                kind = "transactions"
            elif months:
                is_opening = any("остат" in _norm(getattr(c, "value", c)) for r in preview[index+1:index+3] for c in r)
                kind = "stock" if is_opening or preferred == "stock" else "sales"
            elif cols["moq"] is not None or cols["pack"] is not None:
                kind = "moq"
            elif cols["month"] is not None and cols["stock"] is not None:
                kind = "stock"
            elif cols["month"] is not None and cols["qty"] is not None:
                kind = "sales"
        if cols["month"] is not None and cols["factor"] is not None:
            kind = "seasonality"
        if cols["growth"] is not None and cols["category"] is not None:
            kind = "growth"
        if kind:
            score = (100 if kind == preferred else 0) + len(months) + (10 if cols["sku"] is not None else 0)
            candidates.append((score, index, kind, cols, months, transit_columns))
    if not candidates:
        return None
    result = max(candidates, key=lambda item: item[0])
    # Do not accidentally consume duplicate analytical sheets embedded in the
    # monthly and transit workbooks. Separate uploads provide those sources.
    if preferred and result[2] != preferred:
        return None
    return result[1:]


def _eta_from_header(header: Any, filename: str) -> pd.Timestamp | None:
    text = _text(header)
    # A bare "in transit 24.09" may label a snapshot, not an arrival promise.
    if not re.search(r"поступлен|прибыти|arrival|eta", text, re.I):
        return None
    dates = re.findall(r"\b\d{1,2}\.\d{1,2}\.20\d{2}\b", text)
    if dates:
        return _date(dates[-1])
    short = re.search(r"\b(\d{1,2})\.(\d{1,2})\b", text)
    year = re.search(r"20\d{2}", filename)
    if short and year:
        try:
            return pd.Timestamp(int(year.group()), int(short[2]), int(short[1]))
        except ValueError:
            pass
    return None


def _empty() -> dict:
    return {**{key: pd.DataFrame(columns=columns) for key, columns in SCHEMAS.items()}, "warnings": [], "sources": []}


def load_partner_files(files: dict[str, bytes | Path], supplier: str) -> dict:
    """Load uploaded bytes or local paths without copying or publishing originals.

    Returns seven DataFrames, ``warnings`` and ``sources``. SKU keys use internal
    1C codes, when present. Sparse sales cells mean zero. Blank inventory cells
    remain unknown (NaN) and must not silently be interpreted as zero on hand.
    Unknown customer identities stay blank; available identities get a random,
    per-import salted hash. Exceptions are sanitized to avoid source contents.
    """
    result = _empty()
    rows_by_kind = {key: [] for key in SCHEMAS}
    warnings = result["warnings"]
    categories: dict[str, str] = {}
    names: dict[str, str] = {}
    source_signatures = set()
    customer_salt = secrets.token_bytes(32)
    # Cache repeated document timestamps rather than parsing each SKU line again.
    dates_cache: dict[str, pd.Timestamp | None] = {}
    for filename, value in files.items():
        filename = Path(filename).name
        wb = None
        try:
            payload = value if isinstance(value, bytes) else Path(value).read_bytes()
            fingerprint = sha256(payload).digest()
            if fingerprint in source_signatures:
                warnings.append(f"{filename}: повторная копия файла пропущена.")
                continue
            source_signatures.add(fingerprint)
            wb = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            preferred = _file_kind(filename)
            accepted = False
            for ws in wb.worksheets:
                iterator = ws.iter_rows()
                preview = list(islice(iterator, 30))
                detected = _detect(preview, preferred)
                if detected is None:
                    continue
                accepted = True
                head, kind, cols, month_cols, transit_cols = detected
                local_rows = []
                issues = Counter()
                seen_rows = 0
                blank_months = 0
                opening_stock = any("нач. остаток" in _norm(c.value) for r in preview for c in r)
                stock_basis = "opening" if opening_stock else "unknown"
                snapshot_match = re.search(r"на\s+(\d{1,2}\.\d{1,2}\.20\d{2})", filename, re.I)
                snapshot_date = _date(snapshot_match[1]) if snapshot_match else None
                raw_headers = [_norm(c.value) for c in preview[head]]
                free_stock_col = _find(raw_headers, ("свободный остаток",))
                for cells in chain(preview[head+1:], iterator):
                    values = [cell.value for cell in cells]
                    def get(key):
                        idx = cols.get(key)
                        return values[idx] if idx is not None and idx < len(values) else None
                    if not any(v is not None for v in values):
                        continue
                    if kind == "seasonality":
                        if local_rows and _norm(get("month")) in {"итого", "total", "всего"}:
                            break
                        month = _month_number(get("month"))
                        if month is None:
                            continue
                        factor = _number(get("factor"))
                        if factor is None or factor <= 0:
                            issues["неверный коэффициент сезонности"] += 1
                            continue
                        local_rows.append((supplier, _text(get("category")) or ALL_CATEGORY, month, factor))
                        continue
                    if kind == "growth":
                        rate = _number(get("growth"))
                        if rate is None or rate < -1:
                            issues["неверный темп роста"] += 1
                            continue
                        local_rows.append((supplier, _text(get("category")) or ALL_CATEGORY, rate))
                        continue
                    idx = cols["sku"]
                    sku = _sku(cells[idx]) if idx is not None and idx < len(cells) else ""
                    name = _text(get("name"))
                    if not sku or re.match(r"^(итого|всего|total)(?:\s|$)", _norm(name)):
                        continue
                    seen_rows += 1
                    if name:
                        names[sku] = name
                    if get("category") is not None:
                        categories[sku] = _text(get("category")) or ALL_CATEGORY
                    if kind in ("sales", "stock"):
                        if month_cols:
                            for column, month in month_cols.items():
                                raw = values[column] if column < len(values) else None
                                if raw is None or _text(raw) == "":
                                    qty = 0.0 if kind == "sales" else float("nan")
                                    blank_months += 1
                                else:
                                    qty = _number(raw)
                                if qty is None:
                                    issues["нечисловое количество"] += 1
                                    if kind == "sales":
                                        continue
                                    qty = float("nan")
                                if kind == "sales":
                                    local_rows.append((supplier, sku, name, categories.get(sku, ALL_CATEGORY), month, qty))
                                else:
                                    local_rows.append((supplier, sku, month, qty, stock_basis))
                        else:
                            month = _month(get("month")) or _date(get("month"))
                            qty = _number(get("qty" if kind == "sales" else "stock"))
                            if month is None or qty is None:
                                issues["неверный месяц или количество"] += 1
                                continue
                            month = month.to_period("M").to_timestamp()
                            if kind == "sales":
                                local_rows.append((supplier, sku, name, categories.get(sku, ALL_CATEGORY), month, qty))
                            else:
                                local_rows.append((supplier, sku, month, qty, stock_basis))
                    elif kind == "moq":
                        minimum = _number(get("moq"))
                        pack = _number(get("pack"))
                        if minimum is None:
                            minimum = pack
                        if minimum is None or minimum <= 0 or (pack is not None and pack <= 0):
                            issues["неверный MOQ/кратность"] += 1
                            continue
                        local_rows.append((supplier, sku, minimum, pack or 1.0))
                    elif kind == "transit":
                        if free_stock_col is not None and snapshot_date is not None:
                            free_stock = _number(values[free_stock_col])
                            rows_by_kind["stock"].append((supplier, sku, snapshot_date, free_stock if free_stock is not None else float("nan"), "snapshot"))
                        if transit_cols:
                            quantities = [(values[c] if c < len(values) else None, _eta_from_header(h, filename)) for c, h in transit_cols.items()]
                        else:
                            quantities = [(get("qty"), _date(get("eta")))]
                        for raw, eta in quantities:
                            if raw is None or _text(raw) == "":
                                continue
                            qty = _number(raw)
                            if qty is None or qty < 0:
                                issues["неверное количество в пути"] += 1
                                continue
                            if qty:
                                local_rows.append((supplier, sku, qty, eta or pd.NaT))
                                if eta is None:
                                    issues["нет даты поступления"] += 1
                    elif kind == "transactions":
                        raw_date = get("date")
                        cache_key = _text(raw_date)
                        if cache_key not in dates_cache:
                            dates_cache[cache_key] = _date(raw_date)
                        timestamp = dates_cache[cache_key]
                        qty = _number(get("qty"))
                        if timestamp is None or qty is None:
                            issues["неверная дата или количество операции"] += 1
                            continue
                        customer = _text(get("customer"))
                        customer_id = sha256(customer_salt + customer.encode("utf-8")).hexdigest()[:20] if customer else ""
                        local_rows.append((supplier, sku, timestamp, qty, customer_id))
                rows_by_kind[kind].extend(local_rows)
                result["sources"].append({"file": filename, "sheet": ws.title, "kind": kind, "rows": len(local_rows)})
                for issue, count in issues.items():
                    action = "количества сохранены, ETA неизвестна" if issue == "нет даты поступления" else "некорректные значения пропущены"
                    warnings.append(f"{filename}: {issue} — {count}; {action}.")
                if blank_months:
                    if kind == "sales":
                        warnings.append(f"{filename}: пустые месячные ячейки продаж ({blank_months}) приняты за ноль согласно разреженной выгрузке 1С.")
                    else:
                        warnings.append(f"{filename}: пустые остатки ({blank_months}) оставлены неизвестными; требуется проверка склада.")
                if kind == "stock" and opening_stock:
                    warnings.append(f"{filename}: остатки указаны на начало месяца, а не на дату расчёта; обновите склад перед отправкой заказа.")
                elif kind == "stock":
                    warnings.append(f"{filename}: дата внутри месяца для остатка не указана (stock_basis=unknown).")
                if kind == "transactions" and cols["customer"] is None:
                    warnings.append(f"{filename}: идентификаторов покупателей нет; доступен анализ операций/SKU, но не разовых клиентов.")
                if kind == "transit" and any(re.search(r"\b\d{1,2}\.\d{1,2}\b", _text(h)) and not re.search(r"20\d{2}", _text(h)) for h in transit_cols.values()):
                    warnings.append(f"{filename}: краткая дата возле «в пути» не подтверждает дату прибытия; ETA оставлена неизвестной, уточните у поставщика.")
                if kind == "transit" and free_stock_col is not None and snapshot_date is not None:
                    warnings.append(f"{filename}: «Свободный остаток» загружен как снимок на {snapshot_date:%d.%m.%Y}; дата взята из имени отчёта.")
                if kind == "moq" and cols["pack"] is None:
                    warnings.append(f"{filename}: указан минимум отгрузки; отдельная кратность отсутствует и принята равной 1.")
            if not accepted:
                warnings.append(f"{filename}: поддерживаемая таблица не найдена. Нужны код SKU и заголовки показателей.")
        except Exception as exc:
            # User-supplied filenames remain visible; exception text may include
            # private cell contents, so only its class is exposed.
            warnings.append(f"{filename}: файл не прочитан ({type(exc).__name__}). Проверьте формат XLSX и отсутствие пароля.")
        finally:
            if wb is not None:
                wb.close()

    for key, columns in SCHEMAS.items():
        result[key] = pd.DataFrame(rows_by_kind[key], columns=columns)
    sales = result["sales"]
    if sales.empty and not result["transactions"].empty:
        tx = result["transactions"].copy()
        tx["month"] = tx["date"].dt.to_period("M").dt.to_timestamp()
        sales = tx.groupby(["supplier", "sku", "month"], as_index=False)["qty"].sum()
        sales["name"] = sales["sku"].map(names).fillna(sales["sku"])
        sales["category"] = sales["sku"].map(categories).fillna(ALL_CATEGORY)
        result["sales"] = sales[SCHEMAS["sales"]]
        warnings.append("Ежемесячный файл продаж отсутствует: продажи агрегированы из журнала операций.")
    elif not sales.empty:
        result["sales"]["category"] = sales["sku"].map(categories).fillna(sales["category"])
    if not categories and not sales.empty:
        warnings.append("В исходных данных нет товарных категорий: используется общая группа «Все».")
    if not result["transactions"].empty:
        calculated_growth = _journal_growth(result["transactions"], categories)
        if not calculated_growth.empty:
            result["growth"] = calculated_growth
            warnings.append("Рост рассчитан по чистым количествам из журнала: одинаковые месяцы двух последних лет, последний месяц исключён как потенциально неполный. Журнал не суммируется с месячными продажами.")
    for key in ("sales", "stock", "moq", "transit", "seasonality", "transactions"):
        if result[key].empty:
            warnings.append(f"{supplier}: нет данных «{key}»; проверьте комплект исходных файлов.")
    if not result["seasonality"].empty:
        warnings.append("Сезонность применяется на уровне поставщика/указанной группы; денежная сезонность из исходного отчёта является приближением для спроса в штуках.")
    return _deduplicate(result)


def _journal_growth(transactions: pd.DataFrame, categories: dict[str, str]) -> pd.DataFrame:
    tx = transactions.copy()
    tx["date"] = pd.to_datetime(tx["date"])
    records = []
    for supplier, group in tx.groupby("supplier"):
        last = group["date"].max()
        end_month = last.month - 1
        year = last.year
        if end_month == 0:
            year -= 1
            end_month = 12
        group = group[group["date"].dt.month <= end_month].copy()
        group["year"] = group["date"].dt.year
        group["category"] = group["sku"].map(categories).fillna(ALL_CATEGORY)
        # Keep a supplier-wide fallback even when only part of the range has an
        # explicit category in the planning workbook.
        all_groups = list(group.groupby("category"))
        if categories:
            all_groups = [(category, g) for category, g in all_groups if category != ALL_CATEGORY]
            all_groups.append((ALL_CATEGORY, group))
        for category, cat in all_groups:
            old = cat[cat["year"] == year - 1]
            new = cat[cat["year"] == year]
            if old.empty or new.empty:
                continue
            # Compare only calendar months present in both years. Missing
            # transaction months may mean an incomplete source, not zero sales.
            shared = set(old["date"].dt.month) & set(new["date"].dt.month)
            old_total = old.loc[old["date"].dt.month.isin(shared), "qty"].sum()
            new_total = new.loc[new["date"].dt.month.isin(shared), "qty"].sum()
            if old_total > 0:
                records.append((supplier, category, max(0.0, new_total) / old_total - 1.0))
    return pd.DataFrame(records, columns=SCHEMAS["growth"])


def _deduplicate(dataset: dict) -> dict:
    keys = {
        "sales": ["supplier", "sku", "month"], "stock": ["supplier", "sku", "month"],
        "moq": ["supplier", "sku"], "seasonality": ["supplier", "category", "month"],
        "growth": ["supplier", "category"],
    }
    for kind, subset in keys.items():
        frame = dataset[kind]
        duplicates = frame.duplicated(subset, keep="first")
        if duplicates.any():
            dataset["warnings"].append(f"{kind}: повторяющихся ключей {int(duplicates.sum())}; сохранено первое значение, источники не суммировались.")
            dataset[kind] = frame.loc[~duplicates].reset_index(drop=True)
    dataset["warnings"] = list(dict.fromkeys(dataset["warnings"]))
    return dataset


def combine_datasets(datasets: list[dict]) -> dict:
    """Combine supplier imports; supplier is always part of every matching key."""
    result = _empty()
    for key, columns in SCHEMAS.items():
        frames = [d[key] for d in datasets if key in d and not d[key].empty]
        result[key] = pd.concat(frames, ignore_index=True).reindex(columns=columns) if frames else pd.DataFrame(columns=columns)
    for dataset in datasets:
        result["warnings"].extend(dataset.get("warnings", []))
        result["sources"].extend(dataset.get("sources", []))
    return _deduplicate(result)
